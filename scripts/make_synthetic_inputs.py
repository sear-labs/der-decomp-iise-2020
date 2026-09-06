#!/usr/bin/env python
"""Fit a generator to the restricted inputs and sample a shippable instance.

    python scripts/make_synthetic_inputs.py --real <indata1.xlsx> --out data/raw/indata-synthetic.xlsx

The model reads exactly two sheets from `indata1.xlsx`:

    Demand_kWh      hourly demand FRACTIONS by timeslice x month, rows 11..754
    CapacityFactor  availability by technology x timeslice x month, rows 11..1499

Both are derived from Pecan Street Dataport (30 houses, Mueller neighbourhood, Austin)
and ERCOT, as the paper states. This script fits aggregate parameters and samples from
them, so the output derives from the parameters rather than from the series. Only
`fitted_parameters.json` is published alongside.

Two properties of the real data are preserved deliberately, both learned the hard way
on the co-optimisation model:

  * **Support.** Values are clamped to each group's observed [min, max]. An unclamped
    lognormal puts capacity factors above 1.0, which is physically meaningless.
  * **The column sum.** Demand fractions sum to 1 per month - the workbook checks this
    on row 8 - so each month is renormalised after sampling.

Values are written at FULL PRECISION. Rounding inputs is not safe in this family of
models; on the co-optimisation model 4-decimal rounding alone made it infeasible.
"""
from __future__ import annotations

import argparse
import json
import math
import random
import statistics as st
from pathlib import Path

import openpyxl

# Row 10 holds month indices and is a HEADER; data starts at row 11. Month indices are
# numeric, so an "all cells numeric" test would accept row 10 as data and overwrite it.
LOW_BAND = 0.009   # the threshold the model itself keys its solar adjustment off

SHEETS = {
    #  sheet            first  last  keycols       months   renormalise per month?
    "Demand_kWh":      (11,  754, (3, 4), range(5, 17), True),
    "CapacityFactor":  (11, 1499, (3, 4), range(5, 17), False),
}


def draw(rng, mean, cv, lo, hi):
    if mean <= 0 or cv <= 0:
        return min(hi, max(lo, mean))
    s2 = math.log(1 + cv * cv)
    v = rng.lognormvariate(math.log(mean) - s2 / 2, math.sqrt(s2))
    return min(hi, max(lo, v))


def fit_and_sample(ws, first, last, keycols, months, renorm, rng):
    rows = []
    for r in range(first, last + 1):
        keys = tuple(ws.cell(r, c).value for c in keycols)
        if any(k is None or str(k).strip() == "" for k in keys):
            continue
        vals = [ws.cell(r, c).value for c in months]
        if all(isinstance(v, (int, float)) for v in vals):
            rows.append((r, keys, vals))
    if not rows:
        raise ValueError(f"{ws.title}: no data rows parsed from {first}..{last}")

    groups: dict[str, list] = {}
    for r, keys, vals in rows:
        groups.setdefault(str(keys[0]), []).append((r, vals))

    fit = {}
    for name, members in groups.items():
        flat = [v for _, vals in members for v in vals]
        nz = [v for v in flat if v > 0]
        # TWO COMPONENTS, not one. Solar is bimodal - roughly half the hours are night.
        # A single lognormal reproduces the MEAN and destroys the shape: measured, it put
        # 1.0% of PV values under 0.009 where the real series has 44.8%. That matters
        # beyond realism, because the model's own commercial-efficiency adjustment keys
        # off a 0.009 threshold, so a profile that never goes dark is lifted at every
        # hour and solar comes out 15% too strong.
        low = [v for v in flat if v <= LOW_BAND]
        high = [v for v in flat if v > LOW_BAND]
        fit[name] = {
            "n_rows": len(members),
            "monthly_mean": [st.mean([vals[i] for _, vals in members]) for i in range(len(list(months)))],
            "low_fraction": len(low) / len(flat) if flat else 0.0,
            "low_mean": st.mean(low) if low else 0.0,
            "low_max": max(low) if low else 0.0,
            "high_mean": st.mean(high) if high else 0.0,
            "high_cv": (st.pstdev(high) / st.mean(high)) if len(high) > 1 and st.mean(high) else 0.0,
            "observed_min": min(flat) if flat else 0.0,
            "observed_max": max(flat) if flat else 0.0,
        }

    mcols = list(months)
    for name, members in groups.items():
        g = fit[name]
        lo, hi = g["observed_min"], g["observed_max"]
        pl = g["low_fraction"]
        # Split the month mean between the two components so the mixture still hits it.
        for i in range(len(mcols)):
            m = g["monthly_mean"][i]
            hi_mean = (m - pl * g["low_mean"]) / max(1e-12, 1 - pl) if pl < 1 else m
            drawn = []
            for _ in members:
                if rng.random() < pl:
                    drawn.append(draw(rng, g["low_mean"], 0.8, lo, g["low_max"] or LOW_BAND))
                else:
                    drawn.append(draw(rng, max(hi_mean, 1e-9), g["high_cv"] or 0.3,
                                      max(lo, LOW_BAND), hi))
            target = g["monthly_mean"][i]
            for _ in range(50):                      # restore the mean clamping removed
                got = sum(drawn) / len(drawn)
                if got <= 0 or abs(got - target) <= 1e-9 + 0.001 * target:
                    break
                head = [j for j, v in enumerate(drawn) if lo < v < hi - 1e-15]
                if not head:
                    break
                k = target / got
                for j in head:
                    drawn[j] = min(hi, max(lo, drawn[j] * k))
            for (r, _), v in zip(members, drawn):
                ws.cell(r, mcols[i]).value = v

    if renorm:
        # Fractions must sum to 1 per month - the workbook asserts this on row 8.
        for i, c in enumerate(mcols):
            tot = sum(ws.cell(r, c).value for r, _, _ in rows)
            if tot > 0:
                for r, _, _ in rows:
                    ws.cell(r, c).value = ws.cell(r, c).value / tot
    return fit, len(rows)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--real", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--params", default=None)
    ap.add_argument("--seed", type=int, default=20260906)
    args = ap.parse_args()

    rng = random.Random(args.seed)
    wb = openpyxl.load_workbook(args.real, data_only=True)   # values, never formulas
    params = {"seed": args.seed}
    for sheet, (first, last, keycols, months, renorm) in SHEETS.items():
        fit, n = fit_and_sample(wb[sheet], first, last, keycols, months, renorm, rng)
        params[sheet] = fit
        print(f"  {sheet}: {n} rows regenerated across {len(fit)} groups")

    for name in list(wb.sheetnames):
        if name not in SHEETS:
            del wb[name]

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    pf = Path(args.params) if args.params else out.with_name("fitted_parameters.json")
    pf.write_text(json.dumps(params, indent=2))
    print(f"  wrote {out} ({out.stat().st_size:,} B) and {pf.name}")
    print(f"  sheets kept: {wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
